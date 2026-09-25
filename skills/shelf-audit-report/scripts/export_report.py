#!/usr/bin/env python3
"""get_store_report outputs (JSON) -> the customer's xlsx: «Данные», «Сводка», «На проверку».

    python export_report.py --input stores.json --out report.xlsx [--draft]

--input: a JSON list of `get_store_report` results (or one result). Only openpyxl is needed.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLS = ["report_id", "report_date", "route_code", "store_code", "store_name", "product_name", "product_subtype",
        "package_format", "shelf_index", "position", "front_facing_count", "status"]
STATUS_RU = {"approved": "утверждено", "edited": "исправлено", "awaiting_review": "на_проверке"}
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def load(path: Path) -> list[dict]:
    d = json.loads(path.read_text(encoding="utf-8"))
    return d if isinstance(d, list) else [d]


def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def autosize(ws):
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, w + 2))


def build(stores: list[dict], draft: bool = False) -> Workbook:
    wb = Workbook()
    data = wb.active
    data.title = "Данные"
    header(data, COLS)
    review = wb.create_sheet("На проверку")
    header(review, COLS + ["review_reason"])
    n_data = 0
    for st in stores:
        for r in st.get("detail", []):
            row = [r["report_id"], r["report_date"], r.get("route_code", ""), r["store_code"], r["store_name"], r["brand"],
                   r.get("subtype") or ("" if r["brand"] == "Tamga" else "не опознано"), r.get("package_format") or "",
                   r["shelf_index"], r.get("position") or "", r["front_facing_count"], STATUS_RU.get(r["status"], r["status"])]
            if r.get("needs_review") and not draft:
                review.append(row + [r.get("review_reason") or ""])
                for c in review[review.max_row]:
                    c.fill = WARN_FILL
                continue
            data.append(row)
            n_data += 1
    last = max(2, data.max_row)
    summ = wb.create_sheet("Сводка")
    summ.append(["СВОДКА — всё считается формулами из листа «Данные»" + (" (ЧЕРНОВИК: включены неутверждённые)" if draft else "")])
    summ.append([])
    summ.append(["store_name", "product_name", "Всего бутылок снаружи", "Кол-во полок с товаром"])
    for c in summ[3]:
        c.font, c.fill = HEAD, HEAD_FILL
    pairs = sorted({(data.cell(i, 5).value, data.cell(i, 6).value) for i in range(2, data.max_row + 1)} - {(None, None)})
    for store, brand in pairs:
        i = summ.max_row + 1
        rng = lambda col: f"Данные!${col}$2:${col}${last}"  # noqa: E731
        summ.append([store, brand, f"=SUMIFS({rng('K')},{rng('E')},A{i},{rng('F')},B{i})",
                     f'=SUMPRODUCT(({rng("E")}=A{i})*({rng("F")}=B{i})/COUNTIFS({rng("E")},{rng("E")}&"",{rng("F")},{rng("F")}&"",{rng("I")},{rng("I")}&""))'])
    summ.append([])
    summ.append(["На проверке (не в итогах), ед.", None, f"=SUM('На проверку'!K2:K{max(2, review.max_row)})"])
    for ws in (data, review, summ):
        autosize(ws)
    return wb


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--draft", action="store_true", help="put needs_review rows into «Данные» too (draft)")
    a = ap.parse_args(argv)
    stores = load(a.input)
    wb = build(stores, a.draft)
    wb.save(a.out)
    n = sum(len(s.get("detail", [])) for s in stores)
    print(f"{a.out}: {n} rows from {len(stores)} store(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
