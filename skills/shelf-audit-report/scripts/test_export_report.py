"""python -m pytest skills/shelf-audit-report/scripts — export layout without LibreOffice."""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from export_report import main  # noqa: E402


def test_export(tmp_path):
    store = {"store_code": "00-1", "detail": [
        {"report_id": "R1", "report_date": "2026-09-25T10:00:00", "store_code": "00-1", "store_name": "Точка", "status": "approved",
         "shelf_index": 1, "position": "центр", "brand": "AI TEA", "subtype": "AI TEA малина", "package_format": "банка 0,45 л",
         "front_facing_count": 4, "needs_review": False},
        {"report_id": "R1", "report_date": "2026-09-25T10:00:00", "store_code": "00-1", "store_name": "Точка", "status": "approved",
         "shelf_index": 2, "position": "слева", "brand": "AI TEA", "subtype": None, "package_format": None,
         "front_facing_count": 2, "needs_review": True, "review_reason": "подвид не опознан"}]}
    inp = tmp_path / "s.json"
    inp.write_text(json.dumps([store], ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "r.xlsx"
    assert main(["--input", str(inp), "--out", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Данные", "На проверку", "Сводка"]
    assert wb["Данные"].max_row == 2 and wb["На проверку"].max_row == 2
    assert wb["На проверку"]["G2"].value == "не опознано"
    assert str(wb["Сводка"]["C4"].value).startswith("=SUMIFS(")
